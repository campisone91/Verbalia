from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Final

import docx
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from config import (
    APP_NAME,
    APP_VERSION,
    DOCX_FONT_NAME,
    DOCX_FONT_SIZE_PT,
    DOCX_MARGIN_CM,
    DOCX_MIN_SIZE_BYTES,
    DOCX_TITLE_MAX_LENGTH,
    DOCX_TRANSCRIPT_FONT_SIZE_PT,
    OTHERS_DISPLAY_NAME,
    OUTPUT_ORGANIZE_BY_DAY,
)
from modules.contracts import (
    DocumentBuildError,
    DocumentValidationError,
    MeetingMetadata,
    TrackLabel,
    replace_with_retry,
)
from modules.logging_setup import get_logger
from modules.transcript_utils import NO_SPEECH_MARKER, parse_transcript_line

logger = get_logger("docx_builder")

WINDOWS_INVALID_CHARS_RE: Final[re.Pattern[str]] = re.compile(r'[<>:"/\\|?*\x00-\x1f\x7f]')
WHITESPACE_RE: Final[re.Pattern[str]] = re.compile(r"\s+")
WINDOWS_RESERVED_NAMES: Final[frozenset[str]] = frozenset(
    {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *(f"COM{i}" for i in range(1, 10)),
        *(f"LPT{i}" for i in range(1, 10)),
    }
)

ACTION_ITEMS_HEADERS: Final[tuple[str, str, str, str]] = (
    "Responsabile",
    "Task / Azione",
    "Priorità",
    "Scadenza / Note",
)
NO_ACTION_ITEMS_TEXT: Final[str] = "Non sono emersi Action Items espliciti."
PART1_HEADING: Final[str] = "BOZZA MAIL DI RECAP"
PART2_HEADING: Final[str] = "VERBALE DETTAGLIATO DELLA RIUNIONE"
PART3_HEADING: Final[str] = "TRASCRIZIONE INTEGRALE"
DAY_FOLDER_RE: Final[re.Pattern[str]] = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DOCX_DATE_PREFIX_RE: Final[re.Pattern[str]] = re.compile(r"^(?P<day>\d{4}-\d{2}-\d{2})_\d{2}-\d{2}_.+\.docx$")

RAW_MARKDOWN_HEADING_RE: Final[re.Pattern[str]] = re.compile(r"^\s*#{1,6}\s+")
RAW_PLACEHOLDER_RE: Final[re.Pattern[str]] = re.compile(r"\{\{[A-Z0-9_]+\}\}")
NUMBERED_LIST_RE: Final[re.Pattern[str]] = re.compile(r"^\s*\d+[.)]\s+(?P<text>.+)$")
BULLET_LIST_RE: Final[re.Pattern[str]] = re.compile(r"^\s*[-*•]\s+(?P<text>.+)$")
TABLE_SEPARATOR_RE: Final[re.Pattern[str]] = re.compile(
    r"^\s*\|?(?:\s*:?-{2,}:?\s*\|)+\s*:?-{2,}:?\s*\|?\s*$"
)
ACTION_ITEMS_SECTION_RE: Final[re.Pattern[str]] = re.compile(
    r"action\s+item|azioni\s+richieste|next\s+step|attività\s+da\s+svolgere|compiti",
    re.IGNORECASE,
)


def sanitize_windows_filename_title(
    title: str,
    max_length: int = DOCX_TITLE_MAX_LENGTH,
) -> str:
    """Sanitizza il titolo della riunione per l'uso in un filename Windows."""
    if not isinstance(title, str):
        return "Meeting"

    cleaned = WINDOWS_INVALID_CHARS_RE.sub(" ", title)
    cleaned = WHITESPACE_RE.sub(" ", cleaned).strip(" .")

    if not cleaned:
        return "Meeting"

    stem_candidate = cleaned.split(".")[0].strip().upper()
    if cleaned.upper() in WINDOWS_RESERVED_NAMES or stem_candidate in WINDOWS_RESERVED_NAMES:
        return "Meeting"

    if max_length > 0 and len(cleaned) > max_length:
        cleaned = cleaned[:max_length].rstrip(" .")

    if not cleaned:
        return "Meeting"

    stem_after = cleaned.split(".")[0].strip().upper()
    if cleaned.upper() in WINDOWS_RESERVED_NAMES or stem_after in WINDOWS_RESERVED_NAMES:
        return "Meeting"

    return cleaned


def day_folder(output_dir: Path, day: datetime) -> Path:
    """Cartella di output del giorno indicato (output/AAAA-MM-GG)."""
    return output_dir / day.strftime("%Y-%m-%d")


def resolve_unique_docx_path(
    output_dir: Path,
    started_at: datetime,
    title: str,
    *,
    organize_by_day: bool | None = None,
) -> Path:
    """Risolve un percorso univoco per il file DOCX senza sovrascrivere file esistenti.

    Con l'organizzazione per giorno il file finisce in output/AAAA-MM-GG/, dove il giorno è
    quello di inizio call; la sottocartella viene creata se manca.
    """
    by_day = OUTPUT_ORGANIZE_BY_DAY if organize_by_day is None else bool(organize_by_day)
    target_dir = day_folder(output_dir, started_at) if by_day else output_dir
    target_dir.mkdir(parents=True, exist_ok=True)

    safe_title = sanitize_windows_filename_title(title)
    timestamp_str = started_at.strftime("%Y-%m-%d_%H-%M")
    base_stem = f"{timestamp_str}_{safe_title}_Recap"

    candidate = target_dir / f"{base_stem}.docx"
    if not candidate.exists():
        return candidate

    counter = 2
    while True:
        candidate = target_dir / f"{base_stem}_{counter}.docx"
        if not candidate.exists():
            return candidate
        counter += 1


def migrate_flat_output_to_daily(output_dir: Path) -> list[tuple[Path, Path]]:
    """Sposta i DOCX rimasti nella radice di output nella cartella del loro giorno (idempotente).

    Il giorno è ricavato dal prefisso AAAA-MM-GG del nome file; i file senza prefisso restano dove sono.
    """
    moved: list[tuple[Path, Path]] = []
    if not output_dir.is_dir():
        return moved
    for entry in sorted(output_dir.iterdir()):
        if not entry.is_file() or entry.suffix.lower() != ".docx":
            continue
        match = _DOCX_DATE_PREFIX_RE.match(entry.name)
        if match is None:
            continue
        target_dir = output_dir / match.group("day")
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / entry.name
        counter = 2
        while target.exists():
            target = target_dir / f"{entry.stem}_{counter}{entry.suffix}"
            counter += 1
        try:
            shutil.move(str(entry), str(target))
        except OSError as exc:
            logger.warning("Migrazione di %s nella cartella del giorno non riuscita: %s", entry.name, exc)
            continue
        moved.append((entry, target))
        logger.info("DOCX migrato nella cartella del giorno: %s -> %s", entry.name, target.parent.name)
    return moved


def latest_output_folder(output_dir: Path, today: datetime | None = None) -> Path:
    """Cartella da aprire con "Apri cartella output": quella di oggi se esiste, altrimenti la radice."""
    reference = today if today is not None else datetime.now()
    candidate = day_folder(output_dir, reference)
    return candidate if candidate.is_dir() else output_dir


def _has_page_break(doc: DocxDocument) -> bool:
    """Verifica la presenza di almeno un'interruzione di pagina esplicita nel documento."""
    body_xml = doc.element.body.xml
    return 'w:type="page"' in body_xml or "w:type='page'" in body_xml


def _contains_raw_markdown(text: str) -> bool:
    """Rileva la presenza di sintassi Markdown grezza o placeholder non sostituiti."""
    if not text:
        return False
    if RAW_MARKDOWN_HEADING_RE.match(text):
        return True
    if "**" in text:
        return True
    if RAW_PLACEHOLDER_RE.search(text):
        return True
    return False


def validate_docx_file(
    docx_path: Path,
    min_size_bytes: int = DOCX_MIN_SIZE_BYTES,
    *,
    require_recap_email: bool = True,
) -> None:
    """Valida l'integrità strutturale e di contenuto del file DOCX generato."""
    if not docx_path.exists() or not docx_path.is_file():
        raise DocumentValidationError(f"Il file DOCX non esiste: {docx_path}")

    file_size = docx_path.stat().st_size
    if file_size < min_size_bytes:
        raise DocumentValidationError(
            f"Dimensione DOCX insufficiente ({file_size} byte < {min_size_bytes} byte): {docx_path}"
        )

    try:
        doc = docx.Document(str(docx_path))
    except Exception as exc:
        raise DocumentValidationError(
            f"Impossibile aprire il file DOCX {docx_path}: {exc}"
        ) from exc

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    if len(paragraphs) < 4:
        raise DocumentValidationError("Il documento DOCX ha un contenuto testuale insufficiente.")

    has_title = any(
        p.style is not None and p.style.name == "Title" and p.text.strip()
        for p in doc.paragraphs
    )
    if not has_title:
        raise DocumentValidationError("Titolo principale mancante nel documento DOCX.")

    part1_idx: int | None = None
    part2_idx: int | None = None
    for idx, text in enumerate(paragraphs):
        if PART1_HEADING in text and part1_idx is None:
            part1_idx = idx
        if PART2_HEADING in text and part2_idx is None:
            part2_idx = idx

    if require_recap_email:
        if part1_idx is None:
            raise DocumentValidationError(f"Sezione obbligatoria mancante: '{PART1_HEADING}'.")
        if part2_idx is None:
            raise DocumentValidationError(f"Sezione obbligatoria mancante: '{PART2_HEADING}'.")
        if part1_idx >= part2_idx:
            raise DocumentValidationError(
                f"Ordine delle sezioni non valido: '{PART1_HEADING}' deve precedere '{PART2_HEADING}'."
            )
        if not _has_page_break(doc):
            raise DocumentValidationError("Page break obbligatorio mancante tra le sezioni.")
    else:
        if part2_idx is None:
            raise DocumentValidationError(f"Sezione obbligatoria mancante: '{PART2_HEADING}'.")

    for p_text in paragraphs:
        if _contains_raw_markdown(p_text):
            raise DocumentValidationError(
                f"Rilevato Markdown grezzo o placeholder non risolto nel testo: {p_text[:80]}"
            )

    action_table_found = False
    expected_headers = list(ACTION_ITEMS_HEADERS)
    for table in doc.tables:
        if not table.rows:
            continue
        header_cells = [cell.text.strip() for cell in table.rows[0].cells]
        if header_cells == expected_headers:
            action_table_found = True
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if _contains_raw_markdown(cell_text):
                    raise DocumentValidationError(
                        "Rilevato Markdown grezzo o placeholder in una cella di tabella."
                    )

    if not action_table_found:
        raise DocumentValidationError(
            "Tabella Action Items con intestazioni standard non trovata nel DOCX."
        )


def _set_cell_shading(cell: _Cell, fill_hex: str) -> None:
    """Imposta il colore di sfondo (shading) di una cella Word."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _append_formatted_runs(paragraph: Paragraph, text: str) -> None:
    """Aggiunge run al paragrafo convertendo il grassetto **testo** e rimuovendo residui grezzi."""
    sanitized = RAW_PLACEHOLDER_RE.sub("", text)
    parts = re.split(r"(\*\*[^*]+\*\*)", sanitized)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            clean_bold = part[2:-2].replace("**", "")
            run = paragraph.add_run(clean_bold)
            run.bold = True
        else:
            clean_normal = part.replace("**", "")
            if clean_normal:
                paragraph.add_run(clean_normal)


def _parse_pipe_row(line: str) -> list[str]:
    """Estrae le celle da una riga di tabella Markdown pipe."""
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip().replace("**", "") for cell in stripped.split("|")]


def _extract_action_items_from_bullet(text: str) -> tuple[str, str, str, str] | None:
    """Estrae una riga (Responsabile, Task, Priorità, Scadenza) da un bullet di Action Item."""
    cleaned = text.strip().replace("**", "")
    if not cleaned or NO_ACTION_ITEMS_TEXT.lower() in cleaned.lower():
        return None
    if cleaned.lower() in {"nessuno", "nessuno.", "nessuna", "nessuna.", "n/a", "nd", "-"}:
        return None

    parts = [p.strip() for p in re.split(r"\s*[|;]\s*", cleaned) if p.strip()]
    if len(parts) >= 4:
        return (parts[0], parts[1], parts[2], parts[3])

    kv_owner = re.search(
        r"(?:responsabile|owner|assegnatario)\s*:\s*([^,\-;|]+)",
        cleaned,
        re.IGNORECASE,
    )
    kv_priority = re.search(
        r"(?:priorit[aà])\s*:\s*([^,\-;|]+)",
        cleaned,
        re.IGNORECASE,
    )
    kv_due = re.search(
        r"(?:scadenza|entro|note|deadline)\s*:\s*(.+)$",
        cleaned,
        re.IGNORECASE,
    )
    if kv_owner or kv_priority or kv_due:
        owner = kv_owner.group(1).strip() if kv_owner else "Da assegnare"
        priority = kv_priority.group(1).strip() if kv_priority else "Media"
        due = kv_due.group(1).strip() if kv_due else "Da definire"
        task_text = cleaned
        for match in (kv_owner, kv_priority, kv_due):
            if match:
                task_text = task_text.replace(match.group(0), "")
        task_text = re.sub(r"^[\s,\-;|/]+|[\s,\-;|/]+$", "", task_text).strip()
        if not task_text:
            task_text = cleaned
        return (owner, task_text, priority, due)

    colon_split = cleaned.split(":", 1)
    if len(colon_split) == 2 and len(colon_split[0].split()) <= 4:
        return (colon_split[0].strip(), colon_split[1].strip(), "Media", "Da definire")

    return ("Da assegnare", cleaned, "Media", "Da definire")


def _sanitize_transcript_text(text: str) -> str:
    """Neutralizza nel parlato trascritto i caratteri che il validatore leggerebbe come Markdown."""
    clean = RAW_PLACEHOLDER_RE.sub("", text).replace("**", "").strip()
    return clean.lstrip("#").strip()


def _clean_agenda_for_docx(agenda: str) -> str:
    from modules.watcher import clean_outlook_agenda

    cleaned = clean_outlook_agenda(agenda or "", max_chars=1500)
    lines = [line.strip().lstrip("#").strip() for line in cleaned.splitlines()]
    return "\n".join(line for line in lines if line).replace("**", "")


class MeetingDocxBuilder:
    """Costruttore del documento Word finale con validazione e scrittura atomica."""

    def build(
        self,
        metadata: MeetingMetadata,
        started_at: datetime,
        ended_at: datetime,
        recap_email_path: Path | None,
        detailed_minutes_path: Path,
        output_dir: Path,
        *,
        transcript_path: Path | None = None,
        user_display_name: str = "",
        organize_by_day: bool | None = None,
    ) -> Path:
        """Costruisce, valida e promuove atomicamente il file DOCX finale.

        Se transcript_path è indicato e il file esiste, il documento include in coda la
        parte 3 con la trascrizione integrale. Se recap_email_path è None, omette la Parte 1.
        """
        has_recap = recap_email_path is not None and recap_email_path.exists()
        recap_text = ""
        if recap_email_path is not None:
            if not recap_email_path.exists():
                raise DocumentBuildError(f"File recap email mancante: {recap_email_path}")
            try:
                recap_text = recap_email_path.read_text(encoding="utf-8").strip()
            except OSError as exc:
                raise DocumentBuildError(f"Errore lettura file REDUCE: {exc}") from exc
            if not recap_text:
                raise DocumentBuildError("Il contenuto della bozza mail di recap è vuoto.")

        if not detailed_minutes_path.exists():
            raise DocumentBuildError(f"File verbale dettagliato mancante: {detailed_minutes_path}")

        try:
            minutes_text = detailed_minutes_path.read_text(encoding="utf-8").strip()
        except OSError as exc:
            raise DocumentBuildError(f"Errore lettura file REDUCE: {exc}") from exc

        if not minutes_text:
            raise DocumentBuildError("Il contenuto del verbale dettagliato è vuoto.")

        transcript_lines: list[str] = []
        if transcript_path is not None and transcript_path.is_file():
            try:
                transcript_lines = transcript_path.read_text(encoding="utf-8").splitlines()
            except (OSError, UnicodeDecodeError) as exc:
                raise DocumentBuildError(f"Errore lettura trascrizione: {exc}") from exc

        output_dir.mkdir(parents=True, exist_ok=True)
        final_path = resolve_unique_docx_path(
            output_dir=output_dir,
            started_at=started_at,
            title=metadata.title,
            organize_by_day=organize_by_day,
        )
        temp_path = final_path.with_suffix(".tmp.docx")

        try:
            doc = docx.Document()
            self._configure_styles_and_layout(doc, metadata)
            self._build_header_and_metadata_box(doc, metadata, started_at, ended_at)

            if has_recap:
                part1_h1 = doc.add_heading(PART1_HEADING, level=1)
                part1_h1.paragraph_format.space_before = Pt(14)
                part1_h1.paragraph_format.space_after = Pt(6)
                self._render_markdown_block(doc, recap_text, is_detailed_minutes=False)
                doc.add_page_break()

            part2_h1 = doc.add_heading(PART2_HEADING, level=1)
            part2_h1.paragraph_format.space_before = Pt(10)
            part2_h1.paragraph_format.space_after = Pt(6)
            self._render_markdown_block(doc, minutes_text, is_detailed_minutes=True)

            if transcript_lines:
                doc.add_page_break()
                part3_h1 = doc.add_heading(PART3_HEADING, level=1)
                part3_h1.paragraph_format.space_before = Pt(10)
                part3_h1.paragraph_format.space_after = Pt(6)
                self._render_transcript(doc, transcript_lines, user_display_name)

            doc.save(str(temp_path))
        except DocumentBuildError:
            raise
        except Exception as exc:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError as cleanup_exc:
                    logger.debug(
                        "Impossibile rimuovere DOCX temporaneo %s: %s",
                        temp_path,
                        cleanup_exc,
                    )
            raise DocumentBuildError(f"Errore durante la generazione del DOCX: {exc}") from exc

        try:
            validate_docx_file(temp_path, require_recap_email=has_recap)
            replace_with_retry(temp_path, final_path)
        except DocumentValidationError:
            raise
        except Exception as exc:
            raise DocumentBuildError(
                f"Errore durante la promozione atomica del DOCX {final_path}: {exc}"
            ) from exc

        logger.info("Documento Word generato e verificato: %s", final_path.name)
        return final_path

    def _configure_styles_and_layout(
        self,
        doc: DocxDocument,
        metadata: MeetingMetadata,
    ) -> None:
        """Configura margini, stili tipografici, core properties e piè di pagina."""
        doc.core_properties.title = metadata.title or "Riunione Teams"
        doc.core_properties.author = APP_NAME
        doc.core_properties.subject = "Verbale e Recap Riunione Microsoft Teams"
        doc.core_properties.comments = (
            f"Generato localmente da {APP_NAME} v{APP_VERSION}. "
            "Elaborazione 100% on-premise tramite faster-whisper e Ollama."
        )

        for section in doc.sections:
            section.top_margin = Cm(DOCX_MARGIN_CM)
            section.bottom_margin = Cm(DOCX_MARGIN_CM)
            section.left_margin = Cm(DOCX_MARGIN_CM)
            section.right_margin = Cm(DOCX_MARGIN_CM)

            footer = section.footer
            footer_para = (
                footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            )
            footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_para.text = ""
            footer_run = footer_para.add_run(
                f"Generato automaticamente in locale da {APP_NAME} v{APP_VERSION} — "
                "Uso interno. Verificare i contenuti prima della diffusione."
            )
            footer_run.font.name = DOCX_FONT_NAME
            footer_run.font.size = Pt(8.5)
            footer_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        normal_style = doc.styles["Normal"]
        normal_style.font.name = DOCX_FONT_NAME
        normal_style.font.size = Pt(DOCX_FONT_SIZE_PT)
        normal_style.paragraph_format.space_after = Pt(4)
        normal_style.paragraph_format.line_spacing = 1.15

        title_style = doc.styles["Title"]
        title_style.font.name = DOCX_FONT_NAME
        title_style.font.size = Pt(20)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        h1_style = doc.styles["Heading 1"]
        h1_style.font.name = DOCX_FONT_NAME
        h1_style.font.size = Pt(15)
        h1_style.font.bold = True
        h1_style.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        h2_style = doc.styles["Heading 2"]
        h2_style.font.name = DOCX_FONT_NAME
        h2_style.font.size = Pt(12)
        h2_style.font.bold = True
        h2_style.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)

        h3_style = doc.styles["Heading 3"]
        h3_style.font.name = DOCX_FONT_NAME
        h3_style.font.size = Pt(11)
        h3_style.font.bold = True
        h3_style.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    def _build_header_and_metadata_box(
        self,
        doc: DocxDocument,
        metadata: MeetingMetadata,
        started_at: datetime,
        ended_at: datetime,
    ) -> None:
        """Crea il titolo del documento e la tabella riepilogativa dei metadati."""
        clean_title = RAW_PLACEHOLDER_RE.sub("", metadata.title or "Riunione Teams").replace(
            "**", ""
        )
        title_para = doc.add_paragraph(style="Title")
        title_para.add_run(clean_title.strip() or "Riunione Teams")
        title_para.paragraph_format.space_after = Pt(8)

        duration_sec = max(0, int((ended_at - started_at).total_seconds()))
        duration_min = max(1, duration_sec // 60) if duration_sec > 0 else 0

        participants_list = list(metadata.required_attendees) + list(
            metadata.optional_attendees
        )
        participants_str = (
            ", ".join(participants_list)
            if participants_list
            else "Non disponibile (rilevamento locale dual-track [IO]/[ALTRI])"
        )

        meta_rows = [
            ("Titolo Riunione", clean_title.strip() or "Riunione Teams"),
            ("Data e Ora", f"{started_at.strftime('%Y-%m-%d %H:%M')} - {ended_at.strftime('%H:%M')}"),
            ("Durata Effettiva", f"{duration_min} minuti ({duration_sec} secondi)"),
            ("Organizzatore", metadata.organizer or "Non disponibile"),
            ("Partecipanti (Calendario)", participants_str),
            ("Fonte Metadati", metadata.source or "fallback"),
        ]
        agenda_text = _clean_agenda_for_docx(metadata.agenda)
        if agenda_text:
            meta_rows.append(("Agenda", agenda_text))

        table = doc.add_table(rows=len(meta_rows), cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for idx, (label, val) in enumerate(meta_rows):
            row = table.rows[idx]
            cell_key = row.cells[0]
            cell_val = row.cells[1]
            _set_cell_shading(cell_key, "EDF2F7")

            p_key = cell_key.paragraphs[0]
            r_key = p_key.add_run(label)
            r_key.bold = True
            r_key.font.name = DOCX_FONT_NAME
            r_key.font.size = Pt(10)

            p_val = cell_val.paragraphs[0]
            clean_val = RAW_PLACEHOLDER_RE.sub("", str(val)).replace("**", "")
            r_val = p_val.add_run(clean_val)
            r_val.font.name = DOCX_FONT_NAME
            r_val.font.size = Pt(10)

    def _render_markdown_block(
        self,
        doc: DocxDocument,
        markdown_text: str,
        *,
        is_detailed_minutes: bool,
    ) -> None:
        """Converte il testo Markdown controllato in elementi nativi python-docx."""
        lines = markdown_text.splitlines()
        idx = 0
        in_action_items_section = False
        action_items_rendered = False
        collected_action_rows: list[tuple[str, str, str, str]] = []

        while idx < len(lines):
            raw_line = lines[idx]
            stripped = raw_line.strip()

            if not stripped:
                idx += 1
                continue

            if stripped in {"---", "***", "___"}:
                idx += 1
                continue

            heading_match = re.match(r"^(?P<hashes>#{1,6})\s+(?P<title>.+)$", stripped)
            if heading_match:
                if in_action_items_section and not action_items_rendered:
                    self._render_action_items_table(doc, collected_action_rows)
                    action_items_rendered = True
                    in_action_items_section = False

                level_hashes = len(heading_match.group("hashes"))
                heading_title = (
                    RAW_PLACEHOLDER_RE.sub("", heading_match.group("title"))
                    .replace("**", "")
                    .strip()
                )
                if heading_title.upper() in {PART1_HEADING, PART2_HEADING}:
                    idx += 1
                    continue

                docx_level = min(4, max(2, level_hashes))
                h_para = doc.add_heading(heading_title, level=docx_level)
                h_para.paragraph_format.space_before = Pt(8)
                h_para.paragraph_format.space_after = Pt(3)

                # Solo le sezioni di secondo livello aprono la tabella Action Items: un sotto-argomento
                # del resoconto intitolato ad es. "Compiti del team" non deve finire in tabella.
                if (
                    is_detailed_minutes
                    and level_hashes <= 2
                    and ACTION_ITEMS_SECTION_RE.search(heading_title)
                ):
                    in_action_items_section = True
                    collected_action_rows = []
                idx += 1
                continue

            if "|" in stripped and (
                stripped.startswith("|") or stripped.count("|") >= 2
            ):
                table_lines: list[str] = []
                while idx < len(lines):
                    candidate = lines[idx].strip()
                    if "|" in candidate and (
                        candidate.startswith("|") or candidate.count("|") >= 2
                    ):
                        table_lines.append(candidate)
                        idx += 1
                    else:
                        break

                parsed_rows: list[list[str]] = []
                for t_line in table_lines:
                    if TABLE_SEPARATOR_RE.match(t_line):
                        continue
                    row_cells = _parse_pipe_row(t_line)
                    if any(cell for cell in row_cells):
                        parsed_rows.append(row_cells)

                if in_action_items_section:
                    for row_idx, r_cells in enumerate(parsed_rows):
                        if row_idx == 0 and any(
                            "responsabile" in c.lower()
                            or "task" in c.lower()
                            or "azione" in c.lower()
                            for c in r_cells
                        ):
                            continue
                        if len(r_cells) == 1 and NO_ACTION_ITEMS_TEXT.lower() in r_cells[0].lower():
                            continue
                        padded = (r_cells + ["", "", "", ""])[:4]
                        if any(padded):
                            collected_action_rows.append(
                                (
                                    padded[0] or "Da assegnare",
                                    padded[1] or "Da verificare",
                                    padded[2] or "Media",
                                    padded[3] or "Da definire",
                                )
                            )
                elif parsed_rows:
                    self._render_generic_pipe_table(doc, parsed_rows)
                continue

            if in_action_items_section:
                bullet_match = BULLET_LIST_RE.match(stripped) or NUMBERED_LIST_RE.match(
                    stripped
                )
                item_text = (
                    bullet_match.group("text") if bullet_match else stripped
                )
                extracted = _extract_action_items_from_bullet(item_text)
                if extracted is not None:
                    collected_action_rows.append(extracted)
                idx += 1
                continue

            bullet_match = BULLET_LIST_RE.match(stripped)
            if bullet_match:
                p_bullet = doc.add_paragraph(style="List Bullet")
                _append_formatted_runs(p_bullet, bullet_match.group("text").strip())
                idx += 1
                continue

            num_match = NUMBERED_LIST_RE.match(stripped)
            if num_match:
                p_num = doc.add_paragraph(style="List Number")
                _append_formatted_runs(p_num, num_match.group("text").strip())
                idx += 1
                continue

            p_norm = doc.add_paragraph()
            _append_formatted_runs(p_norm, stripped)
            idx += 1

        if is_detailed_minutes and not action_items_rendered:
            if not in_action_items_section:
                h_para = doc.add_heading("Action Items", level=2)
                h_para.paragraph_format.space_before = Pt(8)
                h_para.paragraph_format.space_after = Pt(3)
            self._render_action_items_table(doc, collected_action_rows)

    def _render_transcript(
        self,
        doc: DocxDocument,
        transcript_lines: list[str],
        user_display_name: str,
    ) -> None:
        """Rende la trascrizione come [hh:mm:ss] Nome: testo, con font ridotto."""
        io_label = user_display_name.strip() or "Io"
        others_label = OTHERS_DISPLAY_NAME
        rendered = 0
        for raw_line in transcript_lines:
            segment = parse_transcript_line(raw_line)
            if segment is None:
                continue
            seconds = int(segment.start_sec)
            stamp = f"[{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}:{seconds % 60:02d}]"
            if segment.track == TrackLabel.IO:
                speaker = io_label
            elif segment.speaker:
                speaker = segment.speaker
            else:
                speaker = others_label
            text = _sanitize_transcript_text(segment.text)
            if not text:
                continue
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(1)
            para.paragraph_format.line_spacing = 1.0
            stamp_run = para.add_run(f"{stamp} ")
            stamp_run.font.size = Pt(DOCX_TRANSCRIPT_FONT_SIZE_PT)
            stamp_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            speaker_run = para.add_run(f"{speaker}: ")
            speaker_run.bold = True
            speaker_run.font.size = Pt(DOCX_TRANSCRIPT_FONT_SIZE_PT)
            text_run = para.add_run(text)
            text_run.font.size = Pt(DOCX_TRANSCRIPT_FONT_SIZE_PT)
            rendered += 1
        if rendered == 0:
            note = doc.add_paragraph()
            note_run = note.add_run(
                "Nessun parlato rilevato nella trascrizione."
                if any(line.strip() == NO_SPEECH_MARKER for line in transcript_lines)
                else "Trascrizione non disponibile."
            )
            note_run.italic = True

    def _render_generic_pipe_table(
        self,
        doc: DocxDocument,
        rows: list[list[str]],
    ) -> None:
        """Renderizza una tabella pipe generica come tabella Word Table Grid."""
        if not rows:
            return
        col_count = max(len(r) for r in rows)
        if col_count == 0:
            return

        table = doc.add_table(rows=len(rows), cols=col_count)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for r_idx, row_values in enumerate(rows):
            row = table.rows[r_idx]
            for c_idx in range(col_count):
                cell = row.cells[c_idx]
                val = row_values[c_idx] if c_idx < len(row_values) else ""
                clean_val = RAW_PLACEHOLDER_RE.sub("", val).replace("**", "").strip()
                para = cell.paragraphs[0]
                run = para.add_run(clean_val)
                run.font.name = DOCX_FONT_NAME
                run.font.size = Pt(10)
                if r_idx == 0:
                    run.bold = True
                    _set_cell_shading(cell, "EDF2F7")

    def _render_action_items_table(
        self,
        doc: DocxDocument,
        action_rows: list[tuple[str, str, str, str]],
    ) -> None:
        """Renderizza la tabella ufficiale Action Items e la nota di assenza se vuota."""
        table = doc.add_table(rows=1 + len(action_rows), cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        header_row = table.rows[0]
        for col_idx, header_title in enumerate(ACTION_ITEMS_HEADERS):
            cell = header_row.cells[col_idx]
            _set_cell_shading(cell, "1B365D")
            para = cell.paragraphs[0]
            run = para.add_run(header_title)
            run.bold = True
            run.font.name = DOCX_FONT_NAME
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for row_idx, item_tuple in enumerate(action_rows, start=1):
            row = table.rows[row_idx]
            for col_idx, cell_value in enumerate(item_tuple):
                cell = row.cells[col_idx]
                clean_val = (
                    RAW_PLACEHOLDER_RE.sub("", cell_value)
                    .replace("**", "")
                    .lstrip("#")
                    .strip()
                )
                para = cell.paragraphs[0]
                run = para.add_run(clean_val)
                run.font.name = DOCX_FONT_NAME
                run.font.size = Pt(10)

        if not action_rows:
            note_para = doc.add_paragraph()
            note_run = note_para.add_run(NO_ACTION_ITEMS_TEXT)
            note_run.italic = True
