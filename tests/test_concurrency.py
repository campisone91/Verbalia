from __future__ import annotations

import json
import threading
import time
from collections.abc import Callable
from pathlib import Path

import docx
import pytest

from config import EXIT_SUCCESS
from modules.audio_utils import create_silent_wav
from modules.contracts import (
    AudioChunk,
    ChunkStatus,
    MeetingMetadata,
    SessionContext,
    SessionMode,
    ShutdownReason,
    Stage,
    atomic_write_text,
    now_tz,
)
from modules.orchestrator import MeetingScribeApplication
from modules.processing import ProcessingJob

_RECORDERS: list[FakeRecorder] = []


class FakeRecorder:
    """Registratore finto: scrive chunk già normalizzati e li consegna come farebbe FFmpeg."""

    def __init__(
        self,
        *,
        session_id: str,
        session_dir: Path,
        on_chunk_ready: Callable[[AudioChunk], None],
        state_manager: object,
        **_kwargs: object,
    ) -> None:
        self.session_id = session_id
        self.session_dir = session_dir
        self._on_chunk_ready = on_chunk_ready
        self._state_manager = state_manager
        self._finalize = threading.Event()
        self._recording = False
        self._normalized = threading.Event()
        self._next_index = 0
        self.start_thread: threading.Thread | None = None
        self.stop_thread: threading.Thread | None = None
        _RECORDERS.append(self)

    @property
    def is_recording(self) -> bool:
        return self._recording

    @property
    def finalize_event(self) -> threading.Event:
        return self._finalize

    def start(self) -> None:
        self.start_thread = threading.current_thread()
        self._recording = True

    def emit_chunk(self) -> None:
        idx = self._next_index
        self._next_index += 1
        audio_dir = self.session_dir / "audio"
        io_wav = create_silent_wav(audio_dir / f"chunk_{idx:03d}_io_16k.wav", duration_sec=0.5)
        others_wav = create_silent_wav(audio_dir / f"chunk_{idx:03d}_altri_16k.wav", duration_sec=0.5)

        def _mark(c: SessionContext) -> None:
            for target in (c.recorded_chunks, c.normalized_chunks):
                if idx not in target:
                    target.append(idx)

        self._state_manager.mutate(_mark)  # type: ignore[attr-defined]
        self._on_chunk_ready(
            AudioChunk(
                index=idx,
                session_id=self.session_id,
                start_offset_sec=float(idx * 60),
                end_offset_sec=float(idx * 60 + 30),
                io_wav_path=io_wav,
                others_wav_path=others_wav,
                is_final=False,
                created_at=now_tz(),
                status=ChunkStatus.NORMALIZED,
            )
        )

    def stop_capture(self) -> None:
        self.stop_thread = threading.current_thread()
        self._recording = False
        # L'ultimo chunk viene normalizzato in background, come AudioFinalizationWorker
        worker = threading.Thread(target=self._finish_async, daemon=True)
        worker.start()

    def _finish_async(self) -> None:
        time.sleep(0.2)
        self.emit_chunk()
        self._normalized.set()

    def wait_normalized(self, shutdown_event: threading.Event | None = None) -> list[AudioChunk]:
        self._normalized.wait(timeout=10)
        return []


class _SlowSummarizer:
    """Summarizer finto che si ferma nel MAP della prima sessione finché il test non lo sblocca."""

    def __init__(self, map_md: str, minutes_md: str, email_md: str) -> None:
        self.map_md = map_md
        self.minutes_md = minutes_md
        self.email_md = email_md
        self.map_started: dict[str, threading.Event] = {}
        self.release: dict[str, threading.Event] = {}
        self.blocking_title = "Call A"
        self.processed_titles: list[str] = []

    def run_map_phase(self, transcript_path: Path, metadata: MeetingMetadata, session_dir: Path) -> list[Path]:
        self.processed_titles.append(metadata.title)
        started = self.map_started.setdefault(metadata.title, threading.Event())
        started.set()
        if metadata.title == self.blocking_title:
            assert self.release.setdefault(metadata.title, threading.Event()).wait(timeout=20)
        path = session_dir / "maps" / "map_block_000.md"
        atomic_write_text(path, self.map_md)
        return [path]

    def run_reduce_detailed(self, map_paths: object, metadata: MeetingMetadata, session_dir: Path) -> Path:
        path = session_dir / "reduce" / "detailed_minutes.md"
        atomic_write_text(path, self.minutes_md)
        return path

    def run_reduce_email(self, detailed_minutes_path: Path, metadata: MeetingMetadata, session_dir: Path) -> Path:
        path = session_dir / "reduce" / "recap_email.md"
        atomic_write_text(path, self.email_md)
        return path

    def unload_model(self) -> None:
        return None


class _EchoWhisperModel:
    """Trascrive restituendo il nome della sessione letto dal percorso del file audio."""

    def transcribe(self, audio_path: str, **_kwargs: object) -> tuple[list[object], dict[str, str]]:
        session_name = Path(audio_path).parents[1].name
        track = "locale" if "_io_" in Path(audio_path).name else "remoto"

        class _Seg:
            start = 0.1
            end = 0.4
            text = f"Parlato {track} di {session_name}."
            avg_logprob = -0.1

        return ([_Seg()], {"language": "it"})


def _wait_for(predicate: Callable[[], bool], timeout: float = 20.0) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if predicate():
            return True
        time.sleep(0.05)
    return predicate()


@pytest.fixture
def concurrent_app(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    valid_map_markdown: str,
    valid_detailed_minutes_markdown: str,
    valid_recap_email_markdown: str,
) -> tuple[MeetingScribeApplication, _SlowSummarizer]:
    _RECORDERS.clear()
    monkeypatch.setattr(
        "modules.transcriber._default_whisper_model_factory",
        lambda *_a, **_kw: _EchoWhisperModel(),
    )
    monkeypatch.setattr(
        "modules.orchestrator.fetch_outlook_meeting_metadata",
        lambda *_a, **_kw: MeetingMetadata(title="Riunione Teams", source="fallback"),
    )
    app = MeetingScribeApplication(
        headless=True,
        output_dir=tmp_path / "output",
        temp_dir=tmp_path / "temp",
        consent_path=tmp_path / "consent.json",
        recorder_factory=FakeRecorder,
    )
    summarizer = _SlowSummarizer(valid_map_markdown, valid_detailed_minutes_markdown, valid_recap_email_markdown)
    app.summarizer = summarizer  # type: ignore[assignment]
    return app, summarizer


def test_call_b_records_while_call_a_is_in_map(
    concurrent_app: tuple[MeetingScribeApplication, _SlowSummarizer],
    tmp_path: Path,
    _isolated_logging: Path,
) -> None:
    """F4: la call B viene registrata durante il MAP di A; entrambe producono il DOCX senza interferenze."""
    app, summarizer = concurrent_app
    app._start_services()
    app.worker = app._make_worker()
    app.worker.start()
    try:
        assert app.start_live_session(MeetingMetadata(title="Call A", source="manual"), SessionMode.MANUAL)
        session_a = app.live
        assert session_a is not None
        rec_a = _RECORDERS[-1]
        rec_a.emit_chunk()
        app.finalize_live_session(ShutdownReason.USER_REQUEST)
        # Lo standby va ripristinato sullo stesso thread che aveva avviato la cattura
        assert rec_a.stop_thread is rec_a.start_thread
        assert app.live is None

        assert summarizer.map_started.setdefault("Call A", threading.Event()).wait(timeout=20)
        assert app.worker.current_session_id == session_a.session_id

        # Durante il MAP di A parte la call B
        assert app.start_live_session(MeetingMetadata(title="Call B", source="manual"), SessionMode.MANUAL)
        session_b = app.live
        assert session_b is not None
        rec_b = _RECORDERS[-1]
        rec_b.emit_chunk()
        assert _wait_for(lambda: 0 in session_b.state_manager.load().transcribed_chunks)

        # Gli state.json restano isolati: ciascuno con la propria sessione e il proprio stage
        state_a = json.loads((session_a.session_dir / "state.json").read_text(encoding="utf-8"))
        state_b = json.loads((session_b.session_dir / "state.json").read_text(encoding="utf-8"))
        assert state_a["metadata"]["title"] == "Call A"
        assert state_b["metadata"]["title"] == "Call B"
        assert state_a["stage"] == Stage.SUMMARIZING_MAP.value
        assert state_b["stage"] in (Stage.RECORDING.value, Stage.RECORDING_AND_TRANSCRIBING.value)
        assert state_a["errors"] == [] and state_b["errors"] == []

        app.finalize_live_session(ShutdownReason.USER_REQUEST)
        summarizer.release.setdefault("Call A", threading.Event()).set()

        output_dir = tmp_path / "output"
        assert _wait_for(lambda: len(list(output_dir.rglob("*.docx"))) == 2, timeout=30)
        assert _wait_for(lambda: not session_a.session_dir.exists() and not session_b.session_dir.exists())
        assert summarizer.processed_titles == ["Call A", "Call B"]

        # Ogni DOCX contiene in appendice solo il parlato della propria sessione
        for docx_path in output_dir.rglob("*.docx"):
            text = "\n".join(p.text for p in docx.Document(str(docx_path)).paragraphs)
            assert "TRASCRIZIONE INTEGRALE" in text
            own = session_a if "Call A" in docx_path.name else session_b
            other = session_b if own is session_a else session_a
            assert own.session_dir.name in text
            assert other.session_dir.name not in text
            assert docx_path.parent.name == now_tz().strftime("%Y-%m-%d")

        assert app.event_writer is not None
        app.event_writer.flush(timeout=5)
        assert app.registry is not None
        for runtime in (session_a, session_b):
            job = app.registry.get_job(runtime.session_id)
            assert job is not None and job["status"] == "completed"

        # I log dei worker sono attribuiti alla sessione corretta
        log_text = (_isolated_logging / "app.log").read_text(encoding="utf-8")
        for runtime in (session_a, session_b):
            transcribed = [
                line
                for line in log_text.splitlines()
                if "Chunk 000 trascritto" in line and f"session={runtime.session_id}" in line
            ]
            assert transcribed, f"log di trascrizione non attribuito a {runtime.session_id}"
    finally:
        summarizer.release.setdefault("Call A", threading.Event()).set()
        app._perform_coordinated_shutdown()


def test_failed_job_is_retried_then_parked_without_touching_live_session(
    concurrent_app: tuple[MeetingScribeApplication, _SlowSummarizer],
    tmp_path: Path,
) -> None:
    """B3: un errore di elaborazione non ferma la registrazione in corso; retry e poi parcheggio."""
    app, _ = concurrent_app

    class _AlwaysFailing:
        calls = 0

        def run_map_phase(self, *_a: object, **_kw: object) -> list[Path]:
            _AlwaysFailing.calls += 1
            raise RuntimeError("Ollama non disponibile")

    app.summarizer = _AlwaysFailing()  # type: ignore[assignment]
    app._start_services()
    worker = app._make_worker()
    worker._retry_base_delay_sec = 0.05
    app.worker = worker
    worker.start()
    try:
        assert app.start_live_session(MeetingMetadata(title="Call A", source="manual"), SessionMode.MANUAL)
        session_a = app.live
        assert session_a is not None
        _RECORDERS[-1].emit_chunk()
        app.finalize_live_session(ShutdownReason.USER_REQUEST)

        assert app.start_live_session(MeetingMetadata(title="Call B", source="manual"), SessionMode.MANUAL)
        session_b = app.live
        assert session_b is not None
        rec_b = _RECORDERS[-1]

        assert _wait_for(lambda: _AlwaysFailing.calls >= 3, timeout=30)
        assert app.event_writer is not None
        assert _wait_for(
            lambda: app.event_writer.flush(1) and app.registry.get_job(session_a.session_id)["status"] == "parked",  # type: ignore[union-attr,index]
            timeout=10,
        )
        job = app.registry.get_job(session_a.session_id)  # type: ignore[union-attr]
        assert job is not None and job["attempts"] == 3
        assert "Ollama non disponibile" in str(job["error"])

        # La sessione A resta su disco (riprendibile) e B continua a registrare indisturbata
        assert (session_a.session_dir / "state.json").is_file()
        assert session_a.state_manager.load().stage == Stage.ERROR
        assert rec_b.is_recording is True
        assert app.live is session_b
        assert session_b.state_manager.load().stage in (Stage.RECORDING, Stage.RECORDING_AND_TRANSCRIBING)
        assert session_b.state_manager.load().errors == []

        # "Riprova" dal cruscotto rimette in coda la sessione parcheggiata
        assert app.retry_job(session_a.session_id) is True
        assert _wait_for(lambda: _AlwaysFailing.calls >= 4)
    finally:
        app._perform_coordinated_shutdown()


def test_startup_recovery_resumes_crashed_session_including_open_chunk(
    concurrent_app: tuple[MeetingScribeApplication, _SlowSummarizer],
    tmp_path: Path,
    sample_session_context: SessionContext,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """B14: una sessione morta in registrazione viene recuperata all'avvio, compreso il chunk aperto."""
    app, summarizer = concurrent_app
    summarizer.blocking_title = ""
    from modules.state_manager import StateManager

    sm = StateManager(sample_session_context.session_dir, temp_root=tmp_path / "temp")
    sample_session_context.stage = Stage.RECORDING_AND_TRANSCRIBING
    sample_session_context.ended_at = None
    sm.initialize(sample_session_context)
    audio_dir = sample_session_context.session_dir / "audio"
    # Chunk 0 chiuso e dichiarato, chunk 1 ancora aperto al momento del crash (non dichiarato)
    create_silent_wav(audio_dir / "chunk_000_io_raw.wav", duration_sec=0.5)
    create_silent_wav(audio_dir / "chunk_000_altri_raw.wav", duration_sec=0.5)
    create_silent_wav(audio_dir / "chunk_001_io_raw.wav", duration_sec=0.5)
    create_silent_wav(audio_dir / "chunk_001_altri_raw.wav", duration_sec=0.5)

    def _declare_first(c: SessionContext) -> None:
        c.recorded_chunks = [0]

    sm.mutate(_declare_first)

    normalized: list[str] = []

    def _fake_normalize(_ffmpeg: Path, raw_wav: Path, norm_wav: Path, **_kw: object) -> Path:
        create_silent_wav(norm_wav, duration_sec=0.5)
        raw_wav.unlink()
        normalized.append(norm_wav.name)
        return norm_wav

    monkeypatch.setattr("modules.processing.normalize_wav_with_ffmpeg", _fake_normalize)

    app._start_services()
    try:
        assert app._recover_pending_sessions() == 1
        rc = app._run_queued_jobs_synchronously()
        assert rc == EXIT_SUCCESS
        assert "chunk_001_io_16k.wav" in normalized
        assert not sample_session_context.session_dir.exists()
        docx_files = list((tmp_path / "output").rglob("*.docx"))
        assert len(docx_files) == 1
        text = "\n".join(p.text for p in docx.Document(str(docx_files[0])).paragraphs)
        assert text.count(sample_session_context.session_dir.name) >= 2
    finally:
        app._perform_coordinated_shutdown()


def test_recovery_skips_sessions_locked_by_another_owner(
    concurrent_app: tuple[MeetingScribeApplication, _SlowSummarizer],
    tmp_path: Path,
    sample_session_context: SessionContext,
) -> None:
    """La recovery non tocca le sessioni il cui lock è detenuto da un altro processo o worker."""
    app, _ = concurrent_app
    from modules.state_manager import StateManager

    owner = StateManager(sample_session_context.session_dir, temp_root=tmp_path / "temp")
    owner.initialize(sample_session_context)
    owner.acquire_owner_lock()
    try:
        assert app._recover_pending_sessions() == 0
        assert len(app.processing_queue) == 0
    finally:
        owner.release_owner_lock()
    assert app._recover_pending_sessions() == 1
    assert app.processing_queue.contains(sample_session_context.session_id)
    job = app.processing_queue.get(timeout=0.1)
    assert isinstance(job, ProcessingJob)


def test_run_auto_main_loop_handles_dashboard_commands(
    concurrent_app: tuple[MeetingScribeApplication, _SlowSummarizer],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """run_auto: avvio manuale, chiusura call e uscita arrivano dal cruscotto via SQLite."""
    app, summarizer = concurrent_app
    summarizer.blocking_title = ""
    monkeypatch.setattr("modules.watcher.is_teams_process_running", lambda *_a, **_kw: False)
    monkeypatch.setattr("modules.watcher.is_teams_audio_active", lambda *_a, **_kw: False)
    monkeypatch.setattr("modules.watcher.has_teams_call_window_title", lambda *_a, **_kw: False)
    registry = app.registry
    assert registry is not None

    result: list[int] = []
    runner = threading.Thread(target=lambda: result.append(app.run_auto()), daemon=True)
    runner.start()
    try:
        assert _wait_for(lambda: app.watcher is not None, timeout=10)
        registry.enqueue_command("start_manual")
        assert _wait_for(lambda: app.live is not None, timeout=10)
        live = app.live
        assert live is not None
        _RECORDERS[-1].emit_chunk()
        assert _wait_for(lambda: registry.get_status().get("live_session") == live.session_id, timeout=10)

        registry.enqueue_command("finalize_now")
        assert _wait_for(lambda: len(list((tmp_path / "output").rglob("*.docx"))) == 1, timeout=30)
        assert _wait_for(lambda: (registry.get_job(live.session_id) or {}).get("status") == "completed")
    finally:
        registry.enqueue_command("exit")
        runner.join(timeout=30)
    assert not runner.is_alive()
    assert result == [EXIT_SUCCESS]
