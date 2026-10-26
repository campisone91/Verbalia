from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

import docx
import pytest

from modules.contracts import (
    DocumentBuildError,
    DocumentValidationError,
    MeetingMetadata,
)
from modules.docx_builder import (
    ACTION_ITEMS_HEADERS,
    NO_ACTION_ITEMS_TEXT,
    PART1_HEADING,
    PART2_HEADING,
    MeetingDocxBuilder,
    resolve_unique_docx_path,
    sanitize_windows_filename_title,
    validate_docx_file,
)


def test_sanitize_windows_filename_title() -> None:
    """Verifica rimozione caratteri proibiti, punti/spazi finali, lunghezza e nomi riservati."""
    assert sanitize_windows_filename_title('Call <Q4>: "Budget" / Test? * | \\') == "Call Q4 Budget Test"
    assert sanitize_windows_filename_title("Titolo con punti finali...   ") == "Titolo con punti finali"
    assert sanitize_windows_filename_title("   ...   ") == "Meeting"
    assert sanitize_windows_filename_title("") == "Meeting"

    for reserved in ("CON", "prn", "Aux", "NUL", "COM1", "com9", "LPT1", "lpt9", "CON.txt"):
        assert sanitize_windows_filename_title(reserved) == "Meeting"

    long_title = "A" * 120
    sanitized_long = sanitize_windows_filename_title(long_title, max_length=80)
    assert len(sanitized_long) == 80


def test_resolve_unique_docx_path_handles_collisions(
    tmp_path: Path,
    fixed_now: datetime,
) -> None:
    """Verifica la gestione delle collisioni con suffissi _2, _3 senza sovrascrittura."""
    first = resolve_unique_docx_path(tmp_path, fixed_now, "Allineamento Tecnico")
    assert first.name == "2026-10-23_14-00_Allineamento Tecnico_Recap.docx"
    first.write_bytes(b"placeholder")

    second = resolve_unique_docx_path(tmp_path, fixed_now, "Allineamento Tecnico")
    assert second.name == "2026-10-23_14-00_Allineamento Tecnico_Recap_2.docx"
    second.write_bytes(b"placeholder")

    third = resolve_unique_docx_path(tmp_path, fixed_now, "Allineamento Tecnico")
    assert third.name == "2026-10-23_14-00_Allineamento Tecnico_Recap_3.docx"


def test_docx_builder_creates_valid_document_with_both_parts_and_table(
    tmp_path: Path,
    fixed_now: datetime,
    sample_metadata: MeetingMetadata,
    valid_recap_email_markdown: str,
    valid_detailed_minutes_markdown: str,
) -> None:
    """Verifica che MeetingDocxBuilder produca un DOCX valido con entrambe le parti e tabella Action Items."""
    recap_path = tmp_path / "recap_email.md"
    minutes_path = tmp_path / "detailed_minutes.md"
    recap_path.write_text(valid_recap_email_markdown, encoding="utf-8")
    minutes_path.write_text(valid_detailed_minutes_markdown, encoding="utf-8")

    output_dir = tmp_path / "output"
    builder = MeetingDocxBuilder()
    docx_path = builder.build(
        metadata=sample_metadata,
        started_at=fixed_now,
        ended_at=fixed_now + timedelta(minutes=45),
        recap_email_path=recap_path,
        detailed_minutes_path=minutes_path,
        output_dir=output_dir,
    )

    assert docx_path.exists()
    validate_docx_file(docx_path)

    doc = docx.Document(str(docx_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert PART1_HEADING in full_text
    assert PART2_HEADING in full_text
    assert full_text.index(PART1_HEADING) < full_text.index(PART2_HEADING)

    action_tables = [
        t
        for t in doc.tables
        if t.rows and [c.text.strip() for c in t.rows[0].cells] == list(ACTION_ITEMS_HEADERS)
    ]
    assert len(action_tables) == 1
    assert len(action_tables[0].rows) >= 2


def test_docx_builder_handles_empty_action_items(
    tmp_path: Path,
    fixed_now: datetime,
    sample_metadata: MeetingMetadata,
    valid_recap_email_markdown: str,
) -> None:
    """Verifica la presenza della frase standard quando non vi sono Action Items espliciti."""
    recap_path = tmp_path / "recap_email.md"
    minutes_path = tmp_path / "detailed_minutes.md"
    recap_path.write_text(valid_recap_email_markdown, encoding="utf-8")
    minutes_no_actions = (
        "## 1. Executive Summary\nSintesi.\n\n"
        "## 2. Dati della Riunione\n- Titolo: Test\n\n"
        "## 3. Discussione per Argomento\n- Punto 1\n\n"
        "## 4. Decisioni Confermate\n- Nessuna\n\n"
        "## 5. Questioni Aperte\n- Nessuna\n\n"
        "## 6. Rischi\n- Nessuno\n\n"
        "## 7. Dipendenze\n- Nessuna\n\n"
        f"## 8. Action Items\n{NO_ACTION_ITEMS_TEXT}\n"
    )
    minutes_path.write_text(minutes_no_actions, encoding="utf-8")

    builder = MeetingDocxBuilder()
    docx_path = builder.build(
        metadata=sample_metadata,
        started_at=fixed_now,
        ended_at=fixed_now + timedelta(minutes=20),
        recap_email_path=recap_path,
        detailed_minutes_path=minutes_path,
        output_dir=tmp_path / "output",
    )

    validate_docx_file(docx_path)
    doc = docx.Document(str(docx_path))
    all_paragraphs = "\n".join(p.text for p in doc.paragraphs)
    assert NO_ACTION_ITEMS_TEXT in all_paragraphs


def test_validate_docx_file_negative_cases(tmp_path: Path) -> None:
    """Verifica che validate_docx_file rifiuti file mancanti, piccoli, corrotti o incompleti."""
    missing = tmp_path / "non_esiste.docx"
    with pytest.raises(DocumentValidationError, match="non esiste"):
        validate_docx_file(missing)

    small_file = tmp_path / "small.docx"
    small_file.write_bytes(b"12345")
    with pytest.raises(DocumentValidationError, match="insufficiente"):
        validate_docx_file(small_file)

    corrupt_file = tmp_path / "corrupt.docx"
    corrupt_file.write_bytes(b"X" * 5000)
    with pytest.raises(DocumentValidationError, match="Impossibile aprire"):
        validate_docx_file(corrupt_file)

    # Documento valido come ZIP/DOCX ma privo delle sezioni obbligatorie
    incomplete_path = tmp_path / "incomplete.docx"
    doc = docx.Document()
    doc.add_paragraph("Solo un paragrafo", style="Title")
    for i in range(40):
        doc.add_paragraph(f"Riga riempitiva {i} " * 15)
    doc.save(str(incomplete_path))

    with pytest.raises(DocumentValidationError, match="Sezione obbligatoria mancante"):
        validate_docx_file(incomplete_path, min_size_bytes=100)


def test_docx_builder_raises_on_missing_inputs(
    tmp_path: Path,
    fixed_now: datetime,
    sample_metadata: MeetingMetadata,
) -> None:
    """Verifica che build sollevi DocumentBuildError se i file REDUCE non esistono."""
    builder = MeetingDocxBuilder()
    with pytest.raises(DocumentBuildError):
        builder.build(
            metadata=sample_metadata,
            started_at=fixed_now,
            ended_at=fixed_now,
            recap_email_path=tmp_path / "missing_email.md",
            detailed_minutes_path=tmp_path / "missing_minutes.md",
            output_dir=tmp_path / "output",
        )


def test_docx_builder_parses_bullet_action_items_and_generic_tables(
    tmp_path: Path,
    fixed_now: datetime,
    sample_metadata: MeetingMetadata,
    valid_recap_email_markdown: str,
) -> None:
    """Verifica il parsing di Action Items in formato elenco puntato e di tabelle Markdown generiche."""
    recap_path = tmp_path / "recap_email.md"
    minutes_path = tmp_path / "detailed_minutes.md"
    recap_path.write_text(valid_recap_email_markdown, encoding="utf-8")
    minutes_bullet_actions = (
        "## 1. Executive Summary\nSintesi con **grassetto**.\n\n"
        "## 2. Dati della Riunione\n- Titolo: Test\n\n"
        "## 3. Discussione per Argomento\n"
        "### Sottosezione\n"
        "| Parametro | Valore |\n"
        "| --- | --- |\n"
        "| FPS | 60 |\n\n"
        "1. Punto numerato uno\n"
        "2. Punto numerato due\n\n"
        "## 4. Decisioni Confermate\n- Confermato\n\n"
        "## 5. Questioni Aperte\n- Nessuna\n\n"
        "## 6. Rischi\n- Nessuno\n\n"
        "## 7. Dipendenze\n- Nessuna\n\n"
        "## 8. Action Items\n"
        "- Owner: Marco Rossi | Task: Aggiornare schema DB | Scadenza: Venerdì | Priorità: Alta\n"
        "- Inviare report finale — Anna Bianchi — Non specificata — Media\n"
    )
    minutes_path.write_text(minutes_bullet_actions, encoding="utf-8")

    builder = MeetingDocxBuilder()
    docx_path = builder.build(
        metadata=sample_metadata,
        started_at=fixed_now,
        ended_at=fixed_now + timedelta(minutes=25),
        recap_email_path=recap_path,
        detailed_minutes_path=minutes_path,
        output_dir=tmp_path / "output",
    )
    validate_docx_file(docx_path)
    doc = docx.Document(str(docx_path))
    assert len(doc.tables) >= 2


def test_docx_builder_rejects_empty_files_and_parses_action_bullets(
    tmp_path: Path,
    fixed_now: datetime,
    sample_metadata: MeetingMetadata,
) -> None:
    """Verifica rifiuto di file REDUCE vuoti e parsing dei formati di Action Item a bullet."""
    builder = MeetingDocxBuilder()
    output_dir = tmp_path / "output"

    empty_recap = tmp_path / "empty_recap.md"
    empty_recap.write_text("   \n", encoding="utf-8")
    valid_minutes = tmp_path / "minutes.md"
    valid_minutes.write_text("## 1. Executive Summary\nSintesi", encoding="utf-8")

    with pytest.raises(DocumentBuildError, match="bozza mail"):
        builder.build(
            metadata=sample_metadata,
            started_at=fixed_now,
            ended_at=fixed_now,
            recap_email_path=empty_recap,
            detailed_minutes_path=valid_minutes,
            output_dir=output_dir,
        )

    valid_recap = tmp_path / "recap.md"
    valid_recap.write_text("## Oggetto\nTest", encoding="utf-8")
    empty_minutes = tmp_path / "empty_minutes.md"
    empty_minutes.write_text("", encoding="utf-8")

    with pytest.raises(DocumentBuildError, match="verbale dettagliato"):
        builder.build(
            metadata=sample_metadata,
            started_at=fixed_now,
            ended_at=fixed_now,
            recap_email_path=valid_recap,
            detailed_minutes_path=empty_minutes,
            output_dir=output_dir,
        )

    from modules.docx_builder import _extract_action_items_from_bullet
    result = _extract_action_items_from_bullet("Mario Rossi: Configurare server web")
    assert result is not None
    o, t, p, d = result
    assert o == "Mario Rossi"
    assert "Configurare" in t
    assert p == "Media"


