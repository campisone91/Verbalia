from __future__ import annotations

import json
import threading
from dataclasses import replace
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import docx
import pytest

from config import PROMPTS_DIR
from modules.contracts import (
    MeetingMetadata,
    SessionContext,
    ShutdownRequested,
    atomic_write_json,
    atomic_write_text,
)
from modules.docx_builder import (
    PART3_HEADING,
    MeetingDocxBuilder,
    migrate_flat_output_to_daily,
    validate_docx_file,
)
from modules.state_manager import StateManager
from modules.summarizer import (
    OllamaMapReduceSummarizer,
    _validate_map_output,
    assemble_detailed_narrative,
    default_summarizer_options,
    load_and_render_prompt,
    strip_narrative_section,
)

_REDUCE_LLM_OUTPUT = (
    "## Executive Summary\nRiunione su autenticazione e rilascio.\n\n"
    "## Decisioni\n- Adottare RS256.\n\n"
    "## Questioni aperte\n- Rotazione chiavi.\n\n"
    "## Rischi\n- Latenza.\n\n"
    "## Dipendenze\n- Team piattaforma.\n\n"
    "## Action Items\n| Responsabile | Task / Azione | Priorità | Scadenza / Note |\n"
    "| --- | --- | --- | --- |\n| Da assegnare | Benchmark | Alta | Venerdì |\n"
)


class _ScriptedOllama:
    """Finto server Ollama: risponde in base al prompt e registra le richieste."""

    def __init__(self, map_md: str) -> None:
        self.headers: dict[str, str] = {}
        self.map_md = map_md
        self.prompts: list[dict[str, Any]] = []

    def post(self, url: str, json: dict[str, Any], timeout: tuple[int, int], stream: bool = False) -> Any:
        self.prompts.append(json)
        prompt = str(json.get("prompt", ""))
        if "Executive Summary" in prompt and "MATERIALE ESTRATTO" in prompt:
            text = _REDUCE_LLM_OUTPUT
        elif "SINTESI DEL VERBALE" in prompt:
            text = (
                "Oggetto: Recap\nIntroduzione:\nBuongiorno a tutti,\nTL;DR:\n- punto\n"
                "Decisioni:\n- RS256\nNext Steps:\n- Benchmark\nChiusura:\nUn cordiale saluto."
            )
        else:
            text = self.map_md
        return _StreamResponse(text)


class _StreamResponse:
    status_code = 200

    def __init__(self, text: str) -> None:
        self.text = text
        self.closed = False

    def iter_lines(self) -> Any:
        for word in self.text.split(" "):
            yield json.dumps({"response": word + " ", "done": False}).encode()
        yield json.dumps(
            {
                "done": True,
                "done_reason": "stop",
                "eval_count": 77,
                "eval_duration": 7_000_000_000,
                "prompt_eval_count": 500,
                "prompt_eval_duration": 1_000_000_000,
            }
        ).encode()

    def json(self) -> dict[str, Any]:
        return {"done": True, "response": self.text}

    def close(self) -> None:
        self.closed = True


@pytest.fixture
def session(tmp_path: Path, sample_session_context: SessionContext) -> tuple[StateManager, Path]:
    sm = StateManager(sample_session_context.session_dir, temp_root=tmp_path / "temp")
    sm.initialize(sample_session_context)
    transcript = sample_session_context.session_dir / "transcripts" / "merged_transcript.txt"
    lines = [f"[00:{m:02d}:00] [{'IO' if m % 2 else 'ALTRI'}]: Intervento al minuto {m}." for m in range(0, 25)]
    transcript.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return sm, transcript


def test_v2_prompts_render_with_new_placeholders(tmp_path: Path) -> None:
    rendered = load_and_render_prompt(
        PROMPTS_DIR / "map_chunk.txt",
        {
            "MEETING_TITLE": "T",
            "ORGANIZER": "O",
            "PARTICIPANTS": "P",
            "AGENDA": "A",
            "GLOSSARY_SECTION": "",
            "USER_NAME": "Andrea",
            "MEETING_DURATION": "circa 30 minuti",
            "TARGET_WORDS": "circa 300-500 parole",
            "BLOCK_START": "00:00:00",
            "BLOCK_END": "00:10:00",
            "ROLLING_CONTEXT": "-",
            "TRANSCRIPT": "[00:00:01] [IO]: ciao",
        },
    )
    assert "[IO] è Andrea" in rendered
    assert "## Resoconto dettagliato" in rendered
    assert "GLOSSARIO" not in rendered

    # B19: qualsiasi versione intera >= 1 è accettata
    custom = tmp_path / "custom_prompt.txt"
    custom.write_text("PROMPT_VERSION: zero\nPROMPT_NAME: x\nciao", encoding="utf-8")
    with pytest.raises(Exception, match="Versione prompt"):
        load_and_render_prompt(custom, {})
    custom.write_text("PROMPT_VERSION: 7\nPROMPT_NAME: x\nciao", encoding="utf-8")
    assert load_and_render_prompt(custom, {}) == "ciao"


def test_map_blocks_of_ten_minutes_with_streaming_progress_and_stats(
    session: tuple[StateManager, Path],
    sample_metadata: MeetingMetadata,
    sample_session_context: SessionContext,
    valid_map_markdown: str,
) -> None:
    sm, transcript = session
    fake = _ScriptedOllama(valid_map_markdown)
    progress: list[int] = []
    summarizer = OllamaMapReduceSummarizer(
        state_manager=sm,
        http_session=fake,  # type: ignore[arg-type]
        on_progress=progress.append,
        options=replace(default_summarizer_options(), user_name="Andrea", glossary="RS256"),
    )
    map_paths = summarizer.run_map_phase(transcript, sample_metadata, sample_session_context.session_dir)
    # 25 minuti di parlato con blocchi da 10 minuti -> 3 blocchi
    assert len(map_paths) == 3
    assert all(call["stream"] is True for call in fake.prompts)
    assert all(call["options"]["num_predict"] == 3072 for call in fake.prompts)
    assert "Andrea" in fake.prompts[0]["prompt"]
    assert "GLOSSARIO DI DOMINIO" in fake.prompts[0]["prompt"]
    assert progress and progress[-1] == 77
    assert summarizer.last_inference_stats["eval_tokens_per_sec"] == pytest.approx(11.0)
    meta = json.loads(map_paths[2].with_name("map_block_002.meta.json").read_text(encoding="utf-8"))
    assert meta["block_start_sec"] == 1200.0
    assert sm.load().mapped_blocks == [0, 1, 2]

    detailed = summarizer.run_reduce_detailed(
        map_paths,
        sample_metadata,
        sample_session_context.session_dir,
        started_at=sample_session_context.started_at,
        ended_at=sample_session_context.started_at + timedelta(minutes=25),
    )
    text = detailed.read_text(encoding="utf-8")
    # Il resoconto è assemblato dai MAP (senza seconda compressione), con intestazioni temporali
    assert "## 3. Resoconto dettagliato per argomento" in text
    assert text.count("#### Modulo di autenticazione") == 3
    assert "### 00:20:00 – 00:24:00" in text
    assert "## 1. Executive Summary\n\nRiunione su autenticazione" in text
    assert "## 8. Action Items" in text and "| Da assegnare | Benchmark |" in text
    assert "circa 25 minuti" in text
    reduce_calls = [c for c in fake.prompts if "MATERIALE ESTRATTO" in c["prompt"]]
    assert len(reduce_calls) == 1
    assert reduce_calls[0]["options"]["num_predict"] == 4096
    # Il REDUCE riceve le liste, non la narrativa
    assert "Mario apre la discussione" not in reduce_calls[0]["prompt"]

    email = summarizer.run_reduce_email(detailed, sample_metadata, sample_session_context.session_dir)
    email_prompt = fake.prompts[-1]["prompt"]
    assert "Resoconto dettagliato" not in email_prompt
    assert email.is_file()


def test_hierarchical_reduce_when_material_exceeds_budget(
    session: tuple[StateManager, Path],
    sample_metadata: MeetingMetadata,
    sample_session_context: SessionContext,
    valid_map_markdown: str,
) -> None:
    sm, _ = session
    fake = _ScriptedOllama(valid_map_markdown)
    summarizer = OllamaMapReduceSummarizer(state_manager=sm, http_session=fake)  # type: ignore[arg-type]
    maps_dir = sample_session_context.session_dir / "maps"
    long_list = "\n".join(f"- Decisione numero {i} con molti dettagli tecnici rilevanti" for i in range(60))
    map_paths: list[Path] = []
    for idx in range(8):
        path = maps_dir / f"map_block_{idx:03d}.md"
        atomic_write_text(
            path,
            f"## Resoconto dettagliato\n### Tema {idx}\nTesto.\n\n## Decisioni\n{long_list}\n\n"
            "## Action items\n- Nessuna\n\n## Questioni aperte\n- x\n\n## Rischi\n- y\n\n## Rolling Context\nz\n",
        )
        atomic_write_json(
            path.with_name(f"map_block_{idx:03d}.meta.json"),
            {"block_index": idx, "block_start_sec": idx * 600.0, "block_end_sec": (idx + 1) * 600.0},
        )
        map_paths.append(path)
    summarizer.available_reduce_tokens = 2500
    summarizer.run_reduce_detailed(map_paths, sample_metadata, sample_session_context.session_dir)
    reduce_calls = [c for c in fake.prompts if "MATERIALE ESTRATTO" in c["prompt"]]
    assert len(reduce_calls) >= 3
    assert any("Sintesi del gruppo" in c["prompt"] for c in reduce_calls)


def test_shutdown_interrupts_streaming(
    session: tuple[StateManager, Path],
    sample_metadata: MeetingMetadata,
    sample_session_context: SessionContext,
    valid_map_markdown: str,
) -> None:
    sm, transcript = session
    shutdown = threading.Event()

    class _Interrupting(_ScriptedOllama):
        def post(self, url: str, json: dict[str, Any], timeout: tuple[int, int], stream: bool = False) -> Any:
            response = super().post(url, json, timeout, stream)
            shutdown.set()
            return response

    summarizer = OllamaMapReduceSummarizer(
        state_manager=sm,
        http_session=_Interrupting(valid_map_markdown),  # type: ignore[arg-type]
        shutdown_event=shutdown,
    )
    with pytest.raises(ShutdownRequested):
        summarizer.run_map_phase(transcript, sample_metadata, sample_session_context.session_dir)


def test_assemble_narrative_and_strip_for_email(tmp_path: Path, valid_map_markdown: str) -> None:
    path = tmp_path / "maps" / "map_block_001.md"
    atomic_write_text(path, valid_map_markdown)
    narrative = assemble_detailed_narrative([path], block_duration_sec=600)
    assert narrative.startswith("### 00:10:00 – 00:20:00")
    assert "#### Modulo di autenticazione" in narrative
    assert "Rolling Context" not in narrative
    ok, missing = _validate_map_output(valid_map_markdown)
    assert ok, missing

    minutes = "## 1. Executive Summary\nSintesi\n\n## 3. Resoconto dettagliato per argomento\n### x\nlungo\n\n## 4. Decisioni Confermate\n- d\n"
    stripped = strip_narrative_section(minutes)
    assert "lungo" not in stripped and "Decisioni" in stripped and "Sintesi" in stripped


def test_docx_with_transcript_appendix_day_folder_and_clean_agenda(
    tmp_path: Path,
    fixed_now: datetime,
    sample_metadata: MeetingMetadata,
    valid_recap_email_markdown: str,
    valid_detailed_minutes_markdown: str,
) -> None:
    recap = tmp_path / "recap.md"
    minutes = tmp_path / "minutes.md"
    transcript = tmp_path / "merged.txt"
    recap.write_text(valid_recap_email_markdown, encoding="utf-8")
    minutes.write_text(
        valid_detailed_minutes_markdown.replace(
            "## 4. Decisioni Confermate",
            "### 00:00:00 – 00:10:00\n#### Compiti del team\nDettaglio narrativo.\n\n## 4. Decisioni Confermate",
        ),
        encoding="utf-8",
    )
    transcript.write_text(
        "[00:00:05] [IO]: Buongiorno, iniziamo.\n[00:00:09] [ALTRI]: **Confermo** # agenda.\n",
        encoding="utf-8",
    )
    metadata = replace(
        sample_metadata,
        agenda="Punti:\n1. Budget\nMicrosoft Teams meeting\nMeeting ID: 123\nPasscode: segreto",
    )
    output_dir = tmp_path / "output"
    path = MeetingDocxBuilder().build(
        metadata=metadata,
        started_at=fixed_now,
        ended_at=fixed_now + timedelta(minutes=30),
        recap_email_path=recap,
        detailed_minutes_path=minutes,
        output_dir=output_dir,
        transcript_path=transcript,
        user_display_name="Andrea",
    )
    assert path.parent == output_dir / "2026-10-23"
    validate_docx_file(path)
    document = docx.Document(str(path))
    texts = [p.text for p in document.paragraphs]
    assert PART3_HEADING in texts
    assert "[00:00:05] Andrea: Buongiorno, iniziamo." in texts
    assert "[00:00:09] Interlocutori: Confermo # agenda." in texts
    # Un sotto-argomento "Compiti del team" resta nel resoconto e non apre la tabella Action Items
    assert "Dettaglio narrativo." in texts
    headings4 = [p.text for p in document.paragraphs if p.style is not None and p.style.name == "Heading 4"]
    assert "Compiti del team" in headings4
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert any("Budget" in cell for cell in cells)
    assert not any("segreto" in cell or "Meeting ID" in cell for cell in cells)


def test_migrate_flat_output_to_daily_is_idempotent(tmp_path: Path) -> None:
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    (output_dir / "2026-10-26_09-18_Riunione Teams_Recap.docx").write_bytes(b"a")
    (output_dir / "2026-10-25_12-01_Altro_Recap.docx").write_bytes(b"b")
    (output_dir / "note.docx").write_bytes(b"c")
    (output_dir / "2026-10-26").mkdir()
    (output_dir / "2026-10-26" / "2026-10-26_09-18_Riunione Teams_Recap.docx").write_bytes(b"old")
    moved = migrate_flat_output_to_daily(output_dir)
    assert len(moved) == 2
    assert (output_dir / "2026-10-26" / "2026-10-26_09-18_Riunione Teams_Recap_2.docx").read_bytes() == b"a"
    assert (output_dir / "2026-10-25" / "2026-10-25_12-01_Altro_Recap.docx").is_file()
    assert (output_dir / "note.docx").is_file()
    assert migrate_flat_output_to_daily(output_dir) == []
